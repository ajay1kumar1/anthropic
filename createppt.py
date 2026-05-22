from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import requests
import os

# ──────────────────────────────────────────────
# 1. Download the image
# ──────────────────────────────────────────────
image_url = (
    "https://z-cdn-media.chatglm.cn/files/"
    "f611d807-40f9-4911-a141-2c9b0c2c6780.png"
    "?auth_key=1879433415-2d9ecb2540b4446cb3eed3d4e3900311-0-"
    "2e8ab9b76cfbef0de641d3f3f03ab63d"
)
local_image = "AI-Agent-MCP-HLD-Architecture.png"

print("Downloading image …")
resp = requests.get(image_url, timeout=30)
resp.raise_for_status()
with open(local_image, "wb") as f:
    f.write(resp.content)
print(f"Image saved → {local_image}  ({os.path.getsize(local_image):,} bytes)")

# ──────────────────────────────────────────────
# 2. Create the single-slide presentation
# ──────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Use a blank layout
blank_layout = prs.slide_layouts[6]   # layout index 6 = Blank
slide = prs.slides.add_slide(blank_layout)

# ── Background ──
bg = slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0x0D, 0x1B, 0x2A)  # dark navy

# ── Title bar (top) ──
title_left   = Inches(0)
title_top    = Inches(0)
title_width  = prs.slide_width
title_height = Inches(1.0)

txBox = slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Bedrock Agent MCP — High-Level Design Architecture"
p.font.size      = Pt(32)
p.font.bold      = True
p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p.alignment      = PP_ALIGN.CENTER

# ── Subtitle line ──
sub_left   = Inches(0)
sub_top    = Inches(0.85)
sub_width  = prs.slide_width
sub_height = Inches(0.5)

subBox = slide.shapes.add_textbox(sub_left, sub_top, sub_width, sub_height)
tf2 = subBox.text_frame
tf2.word_wrap = True
p2 = tf2.paragraphs[0]
p2.text = ("End-to-End Workflow (Steps 1-9)  ·  AWS Orchestration Layer  "
           "·  MCP Tool Registry  ·  Data & Cross-Cutting Concerns")
p2.font.size      = Pt(14)
p2.font.color.rgb = RGBColor(0x90, 0xCA, 0xF9)   # light blue
p2.alignment      = PP_ALIGN.CENTER

# ── Decorative accent line ──
from pptx.util import Emu
accent = slide.shapes.add_shape(
    1,  # MSO_SHAPE.RECTANGLE
    Inches(4.5), Inches(1.3),
    Inches(4.333), Inches(0.04)
)
accent.fill.solid()
accent.fill.fore_color.rgb = RGBColor(0x00, 0xBC, 0xD4)  # cyan accent
accent.line.fill.background()

# ── Image (centred, preserving aspect ratio) ──
from PIL import Image as PILImage

img = PILImage.open(local_image)
img_w, img_h = img.size          # pixels

# Available area
avail_left   = Inches(0.6)
avail_top    = Inches(1.5)
avail_width  = prs.slide_width  - Inches(1.2)
avail_height = prs.slide_height - Inches(2.3)

# Scale to fit
scale_w = avail_width  / Emu(img_w * 914400 / 96)   # assuming 96 dpi
scale_h = avail_height / Emu(img_h * 914400 / 96)
scale   = min(scale_w, scale_h, 1.0)

final_w = int(Emu(img_w * 914400 / 96) * scale)
final_h = int(Emu(img_h * 914400 / 96) * scale)

# Centre horizontally
img_left = int((prs.slide_width  - final_w) / 2)
img_top  = avail_top

pic = slide.shapes.add_picture(
    local_image,
    img_left, img_top,
    final_w, final_h
)

# ── Footer ──
foot_left   = Inches(0)
foot_top    = prs.slide_height - Inches(0.45)
foot_width  = prs.slide_width
foot_height = Inches(0.4)

footBox = slide.shapes.add_textbox(foot_left, foot_top, foot_width, foot_height)
tf3 = footBox.text_frame
p3 = tf3.paragraphs[0]
p3.text = "Amazon CloudFront  ·  AWS Lambda  ·  Amazon Bedrock Agent  ·  MCP Protocol"
p3.font.size      = Pt(11)
p3.font.color.rgb = RGBColor(0x78, 0x78, 0x78)
p3.alignment      = PP_ALIGN.CENTER

# ──────────────────────────────────────────────
# 3. Save
# ──────────────────────────────────────────────
output_path = "AI_Agent_MCP_Architecture.pptx"
prs.save(output_path)
print(f"\n✅ Presentation saved → {output_path}")